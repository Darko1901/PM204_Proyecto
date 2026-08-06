import io
from datetime import date, datetime, time, timezone

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core import roles
from app.core.database import get_db
from app.core.deps import require_roles
from app.models.catalogo import Producto, Suministro
from app.models.finanzas import Compra, Ticket
from app.models.operacion import Cuenta, DetalleCuenta
from app.schemas.reportes import (
    ProductoVendido,
    ResumenReporte,
    VentaPorDia,
    VentasHoy,
)

router = APIRouter(prefix="/reportes", tags=["reportes"])

# "ventas" es el nombre histórico del reporte de pedidos; se acepta como alias
# de entrada para no romper clientes existentes, pero internamente (título,
# nombre de archivo, hoja de xlsx) todo usa "pedidos".
_TIPOS_VALIDOS = {"pedidos", "ventas", "productos", "inventario"}
_ALIAS_TIPO = {"ventas": "pedidos"}

_LABEL_TIPO_CUENTA = {"en_mesa": "Mesa", "para_llevar": "Para llevar"}


def _normalizar_tipo(tipo: str) -> str:
    return _ALIAS_TIPO.get(tipo, tipo)


def _tipo_normalizado_o_422(tipo: str) -> str:
    if tipo not in _TIPOS_VALIDOS:
        raise HTTPException(status_code=422, detail=f"Tipo de reporte inválido: {tipo}")
    return _normalizar_tipo(tipo)


def _rango_datetime(desde: date | None, hasta: date | None) -> tuple[datetime | None, datetime | None]:
    inicio = datetime.combine(desde, time.min, tzinfo=timezone.utc) if desde else None
    fin = datetime.combine(hasta, time.max, tzinfo=timezone.utc) if hasta else None
    return inicio, fin


def _query_ventas_por_dia(
    db: Session, desde: date | None, hasta: date | None, tipo_cuenta: str | None
) -> list[VentaPorDia]:
    inicio, fin = _rango_datetime(desde, hasta)
    q = (
        select(func.date(Ticket.emitido_en).label("fecha"), func.sum(Ticket.total).label("total"))
        .join(Cuenta, Ticket.cuenta_id == Cuenta.id)
    )
    if inicio:
        q = q.where(Ticket.emitido_en >= inicio)
    if fin:
        q = q.where(Ticket.emitido_en <= fin)
    if tipo_cuenta and tipo_cuenta != "ambos":
        q = q.where(Cuenta.tipo == tipo_cuenta)
    q = q.group_by(func.date(Ticket.emitido_en)).order_by(func.date(Ticket.emitido_en))

    return [VentaPorDia(fecha=row.fecha, total=float(row.total)) for row in db.execute(q).all()]


def _query_productos_vendidos(
    db: Session,
    desde: date | None,
    hasta: date | None,
    categoria: str | None,
    busqueda: str | None = None,
) -> list[tuple[int, str, float]]:
    inicio, fin = _rango_datetime(desde, hasta)
    q = (
        select(
            Producto.id,
            Producto.nombre,
            func.sum(DetalleCuenta.cantidad).label("cantidad"),
        )
        .join(DetalleCuenta, DetalleCuenta.producto_id == Producto.id)
        .join(Cuenta, DetalleCuenta.cuenta_id == Cuenta.id)
        .where(Cuenta.estado == "pagada")
        .where(DetalleCuenta.estado != "cancelado")
    )
    if inicio:
        q = q.where(Cuenta.cerrada_en >= inicio)
    if fin:
        q = q.where(Cuenta.cerrada_en <= fin)
    if categoria:
        q = q.where(Producto.categoria == categoria)
    if busqueda:
        q = q.where(Producto.nombre.ilike(f"%{busqueda}%"))
    q = q.group_by(Producto.id, Producto.nombre)

    filas = db.execute(q).all()
    return [(row.id, row.nombre, float(row.cantidad)) for row in filas]


def _gastos_totales(db: Session, desde: date | None, hasta: date | None) -> float:
    inicio, fin = _rango_datetime(desde, hasta)
    q = select(func.coalesce(func.sum(Compra.total), 0))
    if inicio:
        q = q.where(Compra.comprado_en >= inicio)
    if fin:
        q = q.where(Compra.comprado_en <= fin)
    return float(db.execute(q).scalar_one())


@router.get("/ventas/hoy", response_model=VentasHoy)
def ventas_de_hoy(
    db: Session = Depends(get_db),
    _: object = Depends(require_roles(roles.ADMIN)),
) -> VentasHoy:
    hoy = datetime.now(timezone.utc).date()
    inicio, fin = _rango_datetime(hoy, hoy)
    q = select(func.coalesce(func.sum(Ticket.total), 0), func.count(Ticket.id)).where(
        Ticket.emitido_en >= inicio, Ticket.emitido_en <= fin
    )
    total, count = db.execute(q).one()
    return VentasHoy(total=float(total), tickets_emitidos=count)


@router.get("/resumen", response_model=ResumenReporte)
def resumen(
    desde: date | None = None,
    hasta: date | None = None,
    tipo_cuenta: str | None = None,
    categoria: str | None = None,
    db: Session = Depends(get_db),
    _: object = Depends(require_roles(roles.ADMIN)),
) -> ResumenReporte:
    ventas_por_dia = _query_ventas_por_dia(db, desde, hasta, tipo_cuenta)
    ventas_totales = sum(v.total for v in ventas_por_dia)
    gastos = _gastos_totales(db, desde, hasta)

    productos = _query_productos_vendidos(db, desde, hasta, categoria)
    productos.sort(key=lambda p: p[2], reverse=True)
    top = productos[:5]
    bottom = list(reversed(productos[-5:])) if len(productos) > 5 else list(reversed(productos))

    return ResumenReporte(
        desde=desde,
        hasta=hasta,
        ventas_totales=ventas_totales,
        gastos_totales=gastos,
        ganancias=ventas_totales - gastos,
        ventas_por_dia=ventas_por_dia,
        productos_mas_vendidos=[
            ProductoVendido(producto_id=p[0], nombre=p[1], cantidad=p[2]) for p in top
        ],
        productos_menos_vendidos=[
            ProductoVendido(producto_id=p[0], nombre=p[1], cantidad=p[2]) for p in bottom
        ],
    )


def _texto_filtro_fechas(desde: date | None, hasta: date | None) -> str:
    d = desde.isoformat() if desde else "Todos"
    h = hasta.isoformat() if hasta else "Todos"
    return f"Del {d} al {h}"


def _texto_filtros(
    tipo: str,
    desde: date | None,
    hasta: date | None,
    tipo_cuenta: str | None,
    categoria: str | None,
    solo_bajo_minimo: bool,
    busqueda: str | None = None,
) -> str:
    """Línea legible con los filtros aplicados, para imprimir en PDF/XLSX.
    Solo incluye los filtros relevantes para `tipo` (los mismos que la UI
    muestra); un filtro no enviado se marca explícitamente como "Todos"."""
    if tipo == "pedidos":
        tc = _LABEL_TIPO_CUENTA.get(tipo_cuenta, "Todos")
        return f"{_texto_filtro_fechas(desde, hasta)} · Tipo de cuenta: {tc}"
    if tipo == "productos":
        return (
            f"{_texto_filtro_fechas(desde, hasta)} · Categoría: {categoria or 'Todas'}"
            f" · Búsqueda: {busqueda or 'Todos'}"
        )
    if tipo == "inventario":
        return (
            f"Solo por debajo del mínimo: {'Sí' if solo_bajo_minimo else 'No'}"
            f" · Búsqueda: {busqueda or 'Todos'}"
        )
    return ""


def _filtros_aplicados(
    tipo: str,
    desde: date | None,
    hasta: date | None,
    tipo_cuenta: str | None,
    categoria: str | None,
    solo_bajo_minimo: bool,
    busqueda: str | None = None,
) -> dict:
    if tipo == "pedidos":
        return {
            "desde": desde.isoformat() if desde else None,
            "hasta": hasta.isoformat() if hasta else None,
            "tipo_cuenta": tipo_cuenta or "ambos",
        }
    if tipo == "productos":
        return {
            "desde": desde.isoformat() if desde else None,
            "hasta": hasta.isoformat() if hasta else None,
            "categoria": categoria,
            "busqueda": busqueda,
        }
    if tipo == "inventario":
        return {"solo_bajo_minimo": solo_bajo_minimo, "busqueda": busqueda}
    return {}


def _filas_export(
    db: Session,
    tipo: str,
    desde: date | None,
    hasta: date | None,
    tipo_cuenta: str | None,
    categoria: str | None,
    solo_bajo_minimo: bool = False,
    busqueda: str | None = None,
) -> tuple[list[str], list[list[str]]]:
    if tipo == "pedidos":
        inicio, fin = _rango_datetime(desde, hasta)
        q = select(Ticket, Cuenta).join(Cuenta, Ticket.cuenta_id == Cuenta.id)
        if inicio:
            q = q.where(Ticket.emitido_en >= inicio)
        if fin:
            q = q.where(Ticket.emitido_en <= fin)
        if tipo_cuenta and tipo_cuenta != "ambos":
            q = q.where(Cuenta.tipo == tipo_cuenta)
        q = q.order_by(Ticket.emitido_en)
        filas = db.execute(q).all()
        encabezados = ["Folio", "Fecha", "Tipo de cuenta", "Total"]
        datos = [
            [
                t.folio,
                t.emitido_en.strftime("%Y-%m-%d %H:%M") if t.emitido_en else "",
                # c.tipo es un enum (TipoCuenta); .value evita que openpyxl lo
                # serialice como "TipoCuenta.en_mesa" en vez de "en_mesa".
                c.tipo.value if hasattr(c.tipo, "value") else c.tipo,
                f"{t.total:.2f}",
            ]
            for t, c in filas
        ]
        return encabezados, datos

    if tipo == "productos":
        productos = _query_productos_vendidos(db, desde, hasta, categoria, busqueda)
        productos.sort(key=lambda p: p[2], reverse=True)
        encabezados = ["Producto", "Cantidad vendida"]
        datos = [[nombre, f"{cantidad:g}"] for _id, nombre, cantidad in productos]
        return encabezados, datos

    if tipo == "inventario":
        q = select(Suministro).order_by(Suministro.nombre)
        if solo_bajo_minimo:
            q = q.where(Suministro.stock_actual < Suministro.stock_minimo)
        if busqueda:
            q = q.where(Suministro.nombre.ilike(f"%{busqueda}%"))
        suministros = list(db.execute(q).scalars())
        encabezados = ["Suministro", "Unidad", "Stock actual", "Stock mínimo", "Bajo mínimo"]
        datos = [
            [
                s.nombre,
                s.unidad,
                f"{float(s.stock_actual):g}",
                f"{float(s.stock_minimo):g}",
                "Sí" if float(s.stock_actual) < float(s.stock_minimo) else "No",
            ]
            for s in suministros
        ]
        return encabezados, datos

    raise HTTPException(status_code=422, detail=f"Tipo de reporte inválido: {tipo}")


@router.get("/preview")
def previsualizar_reporte(
    tipo: str = Query(..., description="pedidos | productos | inventario"),
    desde: date | None = None,
    hasta: date | None = None,
    tipo_cuenta: str | None = None,
    categoria: str | None = None,
    solo_bajo_minimo: bool = False,
    busqueda: str | None = Query(None, description="Filtro por nombre (solo productos/inventario)"),
    db: Session = Depends(get_db),
    _: object = Depends(require_roles(roles.ADMIN)),
) -> dict:
    tipo = _tipo_normalizado_o_422(tipo)
    columnas, filas = _filas_export(db, tipo, desde, hasta, tipo_cuenta, categoria, solo_bajo_minimo, busqueda)
    return {
        "columnas": columnas,
        "filas": filas,
        "total_filas": len(filas),
        "filtros_aplicados": _filtros_aplicados(
            tipo, desde, hasta, tipo_cuenta, categoria, solo_bajo_minimo, busqueda
        ),
    }


@router.get("/export/pdf")
def exportar_pdf(
    tipo: str = Query(..., description="pedidos | productos | inventario"),
    desde: date | None = None,
    hasta: date | None = None,
    tipo_cuenta: str | None = None,
    categoria: str | None = None,
    solo_bajo_minimo: bool = False,
    busqueda: str | None = Query(None, description="Filtro por nombre (solo productos/inventario)"),
    db: Session = Depends(get_db),
    _: object = Depends(require_roles(roles.ADMIN)),
) -> StreamingResponse:
    tipo = _tipo_normalizado_o_422(tipo)

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet

    encabezados, datos = _filas_export(db, tipo, desde, hasta, tipo_cuenta, categoria, solo_bajo_minimo, busqueda)
    subtitulo = _texto_filtros(tipo, desde, hasta, tipo_cuenta, categoria, solo_bajo_minimo, busqueda)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    estilos = getSampleStyleSheet()
    titulo = f"Reporte de {tipo} — CoffeeCode Cafetería"
    elementos = [
        Paragraph(titulo, estilos["Title"]),
        Paragraph(subtitulo, estilos["Normal"]),
        Spacer(1, 12),
    ]

    tabla_datos = [encabezados] + datos if datos else [encabezados, ["Sin datos en el rango seleccionado"]]
    tabla = Table(tabla_datos)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2e7d32")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
            ]
        )
    )
    elementos.append(tabla)
    doc.build(elementos)
    buffer.seek(0)

    filename = f"reporte_{tipo}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/xlsx")
def exportar_xlsx(
    tipo: str = Query(..., description="pedidos | productos | inventario"),
    desde: date | None = None,
    hasta: date | None = None,
    tipo_cuenta: str | None = None,
    categoria: str | None = None,
    solo_bajo_minimo: bool = False,
    busqueda: str | None = Query(None, description="Filtro por nombre (solo productos/inventario)"),
    db: Session = Depends(get_db),
    _: object = Depends(require_roles(roles.ADMIN)),
) -> StreamingResponse:
    tipo = _tipo_normalizado_o_422(tipo)

    from openpyxl import Workbook
    from openpyxl.styles import Font

    encabezados, datos = _filas_export(db, tipo, desde, hasta, tipo_cuenta, categoria, solo_bajo_minimo, busqueda)
    subtitulo = _texto_filtros(tipo, desde, hasta, tipo_cuenta, categoria, solo_bajo_minimo, busqueda)

    wb = Workbook()
    ws = wb.active
    ws.title = tipo[:31]

    ws.append([f"Reporte de {tipo} — CoffeeCode Cafetería"])
    ws.append([subtitulo])
    ws.append([])
    ws.append(encabezados)
    fila_encabezado = ws.max_row
    for celda in ws[fila_encabezado]:
        celda.font = Font(bold=True)
    for fila in datos:
        ws.append(fila)
    for columna in ws.columns:
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in columna) + 2
        ws.column_dimensions[columna[0].column_letter].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_{tipo}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
